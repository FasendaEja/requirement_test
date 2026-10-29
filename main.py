# -*- coding: utf-8 -*-
"""
Утилита автоматизированного тестирования требований с помощью ИИ.

Рабочий процесс:
  1. Рекурсивно сканирует папку INPUT_DIR (по умолчанию ./input) и все
     вложенные подпапки, читает файлы требований форматов:
     doc, docx, md, pdf, txt, xlsx, xml, json.
  2. Отправляет содержимое файлов + промт из PROMPT_FILE в ИИ
     (OpenAI-совместимый API, модель qwen3.8, temperature=0.1,
     max_tokens=200000). Запрос синхронный (stream=False): утилита
     дожидается ПОЛНОГО ответа модели и только после этого начинает
     обработку результата.
  3. Сохраняет полученный от ИИ результат в файл вида
     "тестирование требований__ГГГГ-ММ-ДД_ЧЧ-ММ-СС.<расширение>"
     в папке OUTPUT_DIR.
  4. После завершения тестирования формирует два отдельных артефакта:
     - mind-map ("майнд-мапа__<дата_время>.md") — иерархическая карта
       требований/проверок;
     - чек-лист проверок ("чек-лист__<дата_время>.csv") — таблица со
       столбцами: номер, описание проверки, результат (успешно/ошибка/др),
       пример вводимого значения.
     Оба артефакта генерируются отдельными запросами к ИИ (каждый запрос
     дожидается полного ответа, обработка — только после ответа) с
     фолбэком на локальный парсинг отчёта, если ИИ недоступен.

Все параметры вынесены в файл .env (см. ENV_KEYS ниже).
"""

import csv
import json
import os
import re
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv

# Загружаем настройки из .env в директории проекта (не перекрывая уже
# установленные переменные окружения).
BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

# ----------------------------- Настройки из .env ----------------------------

INPUT_DIR = Path(os.getenv("INPUT_DIR", str(BASE_DIR / "input")))
PROMPT_FILE = Path(os.getenv("PROMPT_FILE", str(BASE_DIR / "prompt.txt")))
OUTPUT_DIR = Path(os.getenv("OUTPUT_DIR", str(BASE_DIR / "output")))

# Относительные пути из .env считаем относительно папки проекта.
def _resolve(p: Path) -> Path:
    return p if p.is_absolute() else (BASE_DIR / p).resolve()

INPUT_DIR = _resolve(INPUT_DIR)
PROMPT_FILE = _resolve(PROMPT_FILE)
OUTPUT_DIR = _resolve(OUTPUT_DIR)

SUPPORTED_EXTENSIONS = [
    # Приводим к виду ".doc": нормализуем расширение с точкой в начале.
    ("." + e.strip().lower().lstrip("."))
    for e in os.getenv("SUPPORTED_EXTENSIONS", "doc,docx,md,pdf,txt,xlsx,xml,json").split(",")
    if e.strip()
]

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY", "")
OPENAI_BASE_URL = os.getenv("OPENAI_BASE_URL", "https://api.openai.com/v1")
MODEL = os.getenv("MODEL", "qwen3.8")
TEMPERATURE = float(os.getenv("TEMPERATURE", "0.1"))
MAX_TOKENS = int(os.getenv("MAX_TOKENS", "200000"))
TIMEOUT = float(os.getenv("TIMEOUT", "600"))
MAX_RETRIES = int(os.getenv("MAX_RETRIES", "3"))
RETRY_DELAY = float(os.getenv("RETRY_DELAY", "5"))
SYSTEM_PROMPT = os.getenv("SYSTEM_PROMPT", "Ты — опытный аналитик по тестированию требований.")

RESULT_PREFIX = os.getenv("RESULT_PREFIX", "тестирование требований__")
OUTPUT_EXTENSION = os.getenv("OUTPUT_EXTENSION", ".md")
TIMESTAMP_FORMAT = os.getenv("TIMESTAMP_FORMAT", "%Y-%m-%d_%H-%M-%S")

MAX_FILE_CHARS = int(os.getenv("MAX_FILE_CHARS", "0"))  # 0 = без ограничения

# --- Генерация mind-map и чек-листа после завершения тестирования ---
GENERATE_MINDMAP = os.getenv("GENERATE_MINDMAP", "true").strip().lower() in ("1", "true", "yes", "on", "да")
GENERATE_CHECKLIST = os.getenv("GENERATE_CHECKLIST", "true").strip().lower() in ("1", "true", "yes", "on", "да")
MINDMAP_PROMPT_FILE = _resolve(Path(os.getenv("MINDMAP_PROMPT_FILE", "prompt_mindmap.txt")))
CHECKLIST_PROMPT_FILE = _resolve(Path(os.getenv("CHECKLIST_PROMPT_FILE", "prompt_checklist.txt")))
MINDMAP_PREFIX = os.getenv("MINDMAP_PREFIX", "майнд-мапа__")
MINDMAP_EXTENSION = os.getenv("MINDMAP_EXTENSION", ".md")
CHECKLIST_PREFIX = os.getenv("CHECKLIST_PREFIX", "чек-лист__")
CHECKLIST_EXTENSION = os.getenv("CHECKLIST_EXTENSION", ".csv")


# ------------------------------ Чтение файлов -------------------------------

def read_txt(path: Path) -> str:
    """Чтение простого текста с подбором кодировки."""
    last_err = None
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError as e:
            last_err = e
    return path.read_text(encoding="utf-8", errors="replace")


def read_docx(path: Path) -> str:
    """Чтение .docx через python-docx, при ошибке — через zip+XML."""
    try:
        import docx  # python-docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        text = "\n".join(parts).strip()
        if text:
            return text
    except Exception:
        pass
    # Запасной вариант: читаем document.xml напрямую из zip-архива.
    import zipfile

    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="replace")
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<[^>]+>", "", xml)
    return xml.strip()


def read_pdf(path: Path) -> str:
    """Чтение .pdf через pypdf (с запасным вариантом через PyPDF2)."""
    try:
        from pypdf import PdfReader
    except ImportError:
        from PyPDF2 import PdfReader  # type: ignore
    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n".join(pages).strip()


def read_xlsx(path: Path) -> str:
    """Чтение .xlsx через openpyxl со всех листов в виде TSV."""
    try:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"=== Лист: {ws.title} ===")
            for row in ws.iter_rows(values_only=True):
                cells = [("" if c is None else str(c)) for c in row]
                if any(cells):
                    lines.append("\t".join(cells))
        wb.close()
        return "\n".join(lines).strip()
    except Exception:
        # Запасной вариант: парсим sharedStrings из zip-архива.
        import zipfile

        with zipfile.ZipFile(path) as z:
            try:
                xml = z.read("xl/sharedStrings.xml").decode("utf-8", errors="replace")
            except KeyError:
                return ""
        strings = re.findall(r"<t[^>]*>(.*?)</t>", xml, flags=re.S)
        return "\n".join(strings).strip()


def read_doc(path: Path) -> str:
    """Чтение старого бинарного .doc: antiword -> libreoffice -> fallback."""
    # Вариант 1: antiword
    try:
        res = subprocess.run(
            ["antiword", "-m", "UTF-8.txt", str(path)],
            capture_output=True, timeout=60,
        )
        if res.returncode == 0 and res.stdout.strip():
            return res.stdout.decode("utf-8", errors="replace").strip()
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass
    # Вариант 2: конвертация через libreoffice/soffice в txt
    for exe in ("libreoffice", "soffice"):
        try:
            tmp_dir = BASE_DIR / ".tmp_doc"
            tmp_dir.mkdir(exist_ok=True)
            res = subprocess.run(
                [exe, "--headless", "--convert-to", "txt:Text (encoded):UTF8",
                 "--outdir", str(tmp_dir), str(path)],
                capture_output=True, timeout=120,
            )
            txt_file = tmp_dir / (path.stem + ".txt")
            if res.returncode == 0 and txt_file.exists():
                text = txt_file.read_text(encoding="utf-8", errors="replace").strip()
                txt_file.unlink(missing_ok=True)
                if text:
                    return text
        except (FileNotFoundError, subprocess.TimeoutExpired):
            continue
    # Вариант 3: грубое извлечение печатных последовательностей из бинарника
    data = path.read_bytes()
    words = re.findall(rb"[\x20-\x7e\xd0-\xdf]{4,}", data)
    chunks = []
    for w in words:
        try:
            chunks.append(w.decode("cp1251", errors="ignore"))
        except Exception:
            chunks.append(w.decode("ascii", errors="ignore"))
    return "\n".join(chunks).strip()


def read_xml(path: Path) -> str:
    """Чтение XML как текста с pretty-print, если возможно."""
    raw = read_txt(path)
    try:
        import xml.dom.minidom

        return xml.dom.minidom.parseString(raw).toprettyxml(indent="  ").strip()
    except Exception:
        return raw


def read_json(path: Path) -> str:
    """Чтение JSON как текста с pretty-print, если возможно."""
    raw = read_txt(path)
    try:
        return json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
    except Exception:
        return raw


READERS = {
    ".txt": read_txt,
    ".md": read_txt,
    ".docx": read_docx,
    ".doc": read_doc,
    ".pdf": read_pdf,
    ".xlsx": read_xlsx,
    ".xml": read_xml,
    ".json": read_json,
}


def collect_requirement_files(root: Path) -> list[Path]:
    """Рекурсивный сбор файлов поддерживаемых форматов из root."""
    files = []
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            files.append(path)
    return files


def extract_content(path: Path) -> str:
    """Извлечение текста из файла с учётом его формата."""
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        return ""
    try:
        text = reader(path) or ""
    except Exception as e:
        return f"[Ошибка чтения файла: {e}]"
    if MAX_FILE_CHARS and len(text) > MAX_FILE_CHARS:
        text = text[:MAX_FILE_CHARS] + "\n...[файл обрезано по лимиту MAX_FILE_CHARS]..."
    return text


def build_requirements_payload(files: list[Path], root: Path) -> str:
    """Сбор единого документа из содержимого всех файлов требований."""
    blocks = []
    for i, f in enumerate(files, 1):
        rel = f.relative_to(root)
        content = extract_content(f)
        blocks.append(
            f"===== ФАЙЛ {i}: {rel} =====\n{content}\n===== КОНЕЦ ФАЙЛА {i}: {rel} ====="
        )
    return "\n\n".join(blocks)


# ------------------------------ Работа с ИИ ---------------------------------

def ask_ai_raw(user_message: str, label: str) -> str:
    """
    Низ-level синхронный запрос к ИИ через OpenAI-совместимый API.

    Функция возвращает управление ТОЛЬКО после того, как модель вернёт
    полный ответ (stream=False, обычный blocking-вызов). Обработка
    результата в вызывающем коде начнётся строго после получения ответа.
    Предусмотрены ретраи при сетевых/серверных ошибках.
    """
    from openai import OpenAI  # импорт здесь, чтобы env читался до создания клиента

    client = OpenAI(api_key=OPENAI_API_KEY, base_url=OPENAI_BASE_URL, timeout=TIMEOUT)

    last_error = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            print(f"[{label}: запрос #{attempt}] Отправка данных в ИИ "
                  f"(модель: {MODEL}, temperature={TEMPERATURE}, max_tokens={MAX_TOKENS})...")
            completion = client.chat.completions.create(
                model=MODEL,
                temperature=TEMPERATURE,
                max_tokens=MAX_TOKENS,
                stream=False,  # ждём полный ответ; обработка — только после ответа ИИ
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": user_message},
                ],
            )
            answer = (completion.choices[0].message.content or "").strip()
            print(f"[{label}: ОТВЕТ ПОЛУЧЕН] ИИ вернул полный ответ, начинаю обработку результата.")
            return answer
        except Exception as e:
            last_error = e
            print(f"[{label}: ошибка] Попытка {attempt}/{MAX_RETRIES} не удалась: {e}")
            if attempt < MAX_RETRIES:
                delay = RETRY_DELAY * attempt
                print(f"Повтор через {delay:.0f} с...")
                time.sleep(delay)
    raise RuntimeError(f"Не удалось получить ответ от ИИ ({label}) после {MAX_RETRIES} попыток: {last_error}")


def ask_ai(prompt_text: str, requirements_text: str) -> str:
    """Запрос на тестирование требований: промт из файла + содержимое файлов."""
    user_message = (
        f"{prompt_text}\n\n"
        f"Ниже приведено содержимое файлов с требованиями.\n\n"
        f"{requirements_text}"
    )
    return ask_ai_raw(user_message, "Тестирование требований")


# ------------------------------ Сохранение ----------------------------------

def save_result(answer: str, files_count: int) -> Path:
    """Сохранение ответа ИИ в файл 'тестирование требований__<дата_время>.ext'."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime(TIMESTAMP_FORMAT)
    out_path = OUTPUT_DIR / f"{RESULT_PREFIX}{timestamp}{OUTPUT_EXTENSION}"

    header = (
        f"# Тестирование требований\n\n"
        f"- Дата и время выполнения: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
        f"- Модель: {MODEL}\n"
        f"- temperature={TEMPERATURE}, max_tokens={MAX_TOKENS}\n"
        f"- Файлов с требованиями обработано: {files_count}\n\n"
        f"---\n\n"
    )
    out_path.write_text(header + answer, encoding="utf-8")
    return out_path


# ------------------- Mind-map и чек-лист проверок ---------------------------

CHECKLIST_HEADER = ["Номер", "Описание проверки", "Результат", "Пример вводимого значения"]

# Нормализация вариантов результата к единому виду.
_RESULT_MAP = {
    "успешно": "успешно", "pass": "успешно", "ok": "успешно", "пройдено": "успешно",
    "успех": "успешно", "да": "успешно", "соответствует": "успешно",
    "ошибка": "ошибка", "fail": "ошибка", "failed": "ошибка", "провал": "ошибка",
    "не пройдено": "ошибка", "не соответствует": "ошибка", "нет": "ошибка", "дефект": "ошибка",
    "не применимо": "не применимо", "n/a": "не применимо", "na": "не применимо",
    "не требуется": "не применимо",
    "требуется уточнение": "требуется уточнение", "уточнить": "требуется уточнение",
    "unknown": "требуется уточнение", "неизвестно": "требуется уточнение",
    "?": "требуется уточнение",
}


def normalize_result(value: str) -> str:
    v = (value or "").strip().lower().replace("*", "").replace("`", "")
    if not v:
        return "требуется уточнение"
    if v in _RESULT_MAP:
        return _RESULT_MAP[v]
    for key, norm in _RESULT_MAP.items():
        if key and key in v:
            return norm
    return value.strip()  # нестандартное значение сохраняем как есть


def strip_code_fences(text: str) -> str:
    """Удаляет обрамляющие markdown-блоки ``` ... ``` из ответа ИИ."""
    m = re.search(r"```[a-zA-Z0-9_-]*\s*(.*?)\s*```", text, flags=re.S)
    if m and len(m.group(1).strip()) > 10:
        return m.group(1).strip()
    return text.strip()


def load_aux_prompt(path: Path, default: str) -> str:
    """Читает промт из файла; если файла нет — использует встроенный default."""
    try:
        if path.exists():
            return path.read_text(encoding="utf-8").strip()
    except Exception as e:
        print(f"[Внимание] Не удалось прочитать промт {path}: {e}. Использую встроенный.")
    return default


def ask_ai_or_fallback(user_message: str, label: str, fallback: str) -> str:
    """Запрос к ИИ; при недоступности модели — локальный фолбэк-результат."""
    try:
        return ask_ai_raw(user_message, label)
    except Exception as e:
        print(f"[Внимание] {label}: ИИ недоступен ({e}). Использую локальную генерацию.")
        return fallback


def build_mindmap(report_text: str, requirements_text: str) -> str:
    """Формирует mind-map по результатам тестирования (запрос к ИИ + фолбэк)."""
    default_prompt = (
        "По приведённому отчёту тестирования требований и тексту требований построй "
        "mind-map в формате Markdown-списка с отступами (2 пробела на уровень): "
        "корень — название системы/набора требований; ветви верхнего уровня — темы "
        "(требования, дефекты, риски, вопросы заказчику); далее — конкретные пункты. "
        "Для каждого требования укажи статус проверки (успешно / ошибка / требуется "
        "уточнение). Верни ТОЛЬКО mind-map без пояснительного текста."
    )
    prompt_text = load_aux_prompt(MINDMAP_PROMPT_FILE, default_prompt)
    user_message = (
        f"{prompt_text}\n\n=== ОТЧЁТ ПО РЕЗУЛЬТАТАМ ТЕСТИРОВАНИЯ ТРЕБОВАНИЙ ===\n{report_text}"
        f"\n\n=== ИСХОДНЫЕ ТРЕБОВАНИЯ ===\n{requirements_text}"
    )
    fallback = (
        "# Тестирование требований\n"
        "## Требования\n- см. отчёт (автоматическая декомпозиция недоступна)\n"
        "## Дефекты\n- см. таблицу дефектов отчёта\n"
        "## Вопросы заказчику\n- см. раздел «Вопросы заказчику» отчёта\n"
    )
    return strip_code_fences(ask_ai_or_fallback(user_message, "Mind-map", fallback))


def save_mindmap(mindmap_text: str) -> Path:
    """Сохранение mind-map в отдельный файл 'майнд-мапа__<дата_время>.md'."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime(TIMESTAMP_FORMAT)
    out_path = OUTPUT_DIR / f"{MINDMAP_PREFIX}{timestamp}{MINDMAP_EXTENSION}"
    content = mindmap_text
    if not content.lstrip().startswith("#"):
        content = "# Mind-map: тестирование требований\n\n" + content
    out_path.write_text(content.rstrip() + "\n", encoding="utf-8")
    return out_path


def parse_checklist_rows(text: str) -> list[list[str]]:
    """Парсит строки чек-листа из CSV (';' или ',') либо markdown-таблицы (4 колонки)."""
    rows: list[list[str]] = []
    # Ответ ИИ может быть обёрнут в markdown-блок ``` ... ``` — снимаем его.
    text = strip_code_fences(text)

    def add_row(cells: list[str]):
        if not re.match(r"^\d+", cells[0]):
            return  # пропускаем заголовки и служебные строки
        num = re.sub(r"[^\d].*$", "", cells[0]) or str(len(rows) + 1)
        desc, result, example = cells[1], normalize_result(cells[2]), cells[3]
        if desc:
            rows.append([num, desc, result, example])

    # 1) Markdown-таблица (разбирается первой — её проще детектировать).
    md_rows: list[list[str]] = []
    for line in text.splitlines():
        s = line.strip()
        if s.startswith("|") and s.count("|") >= 3:
            cells = [c.strip() for c in s.strip("|").split("|")]
            if cells and all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue  # строка-разделитель |---|---|
            if len(cells) >= 4:
                md_rows.append(cells[:4])
    for cells in md_rows:
        add_row(cells)
    if rows:
        return rows

    # 2) CSV. Разделитель не угадываем по тексту (кавычки мешают), а
    # перебираем варианты: ';' , ',' , '\t'. Для каждого варианта строки
    # с «неполным» числом колонок склеиваются со следующmи (многострочные
    # ячейки в кавычках csv.reader объединит сам).
    lines = [ln for ln in text.splitlines() if ln.strip()]
    best: list[list[str]] = []
    for delim in (";", ",", "\t"):
        parsed: list[list[str]] = []
        buf = ""
        for ln in lines:
            candidate = (buf + "\n" + ln) if buf else ln
            recs = list(csv.reader(candidate.splitlines(), delimiter=delim))
            if recs and len(recs[0]) >= 4:
                parsed.extend([r for r in recs if len(r) >= 4])
                buf = ""
            else:
                buf = candidate
        scored = [cells[:4] for cells in parsed if len(cells) >= 4]
        if len(scored) > len(best):
            best = scored
    for cells in best:
        add_row(cells)
    return rows


def clean_md(s: str) -> str:
    """Убирает markdown-разметку из строки."""
    s = re.sub(r"\*\*|__|`", "", s or "")
    return s.replace("\n", " ").strip()


def fallback_checklist_rows(report_text: str) -> list[list[str]]:
    """Локальная сборка чек-листа из markdown-таблицы дефектов отчёта."""
    rows: list[list[str]] = []
    for line in report_text.splitlines():
        s = line.strip()
        if not (s.startswith("|") and s.count("|") >= 4):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        # Дефекты описываются таблицей из 6 колонок (№|Файл|Требование|Тип|Описание|Рекомендация).
        if len(cells) < 6 or not re.match(r"^\d+", cells[0]):
            continue
        file_ref, req, problem, description, recommendation = cells[1:6]
        low_problem = problem.lower()
        low_all = " ".join(c.lower() for c in cells)
        if any(k in low_all for k in ("проверить", "валидировать")):
            continue  # это уже строка чек-листа, а не таблица дефектов отчёта
        if any(k in low_problem for k in ("отсутств", "неполн", "противореч", "неоднозначн")):
            result = "ошибка"
        elif any(k in low_problem for k in ("непроверяем", "неясн", "нечётк", "термин", "неопредел")):
            result = "требуется уточнение"
        else:
            result = "не применимо"
        source = file_ref if file_ref and file_ref != "-" else "документ требований"
        desc = (
            f"Проверить требование «{clean_md(req)}» ({source}): {clean_md(description)}. "
            f"Рекомендация: {clean_md(recommendation) or 'см. отчёт'}."
        )
        rows.append([str(len(rows) + 1), desc, result, "см. отчёт тестирования"])
    return rows


def generate_checklist(report_text: str, requirements_text: str) -> tuple[list[list[str]], str]:
    """
    Формирует чек-лист проверок: запрос к ИИ, затем разбор его ответа.
    При недоступности ИИ или нераспознанном формате — локальный фолбэк
    (парсинг таблицы дефектов из отчёта). Возвращает (строки, сырой ответ ИИ).
    """
    default_prompt = (
        "На основе отчёта тестирования требований составь чек-лист проверок. "
        "Верни СТРОГО CSV-таблицу с точкой с запятой (;) как разделителем и заголовком:\n"
        "Номер;Описание проверки;Результат;Пример вводимого значения\n\n"
        "Правила:\n"
        "- Каждое требование и каждый дефект из отчёта — отдельная строка чек-листа.\n"
        "- Колонка «Результат» — одно из: успешно, ошибка, не применимо, требуется уточнение "
        "(требование пройдено / выявлен дефект / проверка неприменима / нужна информация).\n"
        "- В «Пример вводимого значения» указывай конкретные тестовые данные (валидные и "
        "граничные), либо «нет данных», если пример вывести нельзя.\n"
        "- Описания — на русском языке, без переносов строк внутри ячеек.\n"
        "- Никакого текста вне CSV-таблицы."
    )
    prompt_text = load_aux_prompt(CHECKLIST_PROMPT_FILE, default_prompt)
    user_message = (
        f"{prompt_text}\n\n=== ОТЧЁТ ПО РЕЗУЛЬТАТАМ ТЕСТИРОВАНИЯ ТРЕБОВАНИЙ ===\n{report_text}"
        f"\n\n=== ИСХОДНЫЕ ТРЕБОВАНИЯ ===\n{requirements_text}"
    )
    answer = ask_ai_or_fallback(user_message, "Чек-лист", "")
    rows = parse_checklist_rows(answer) if answer else []
    if rows:
        print(f"[Чек-лист] Разобрано строк из ответа ИИ: {len(rows)}")
        return rows, answer
    if answer:
        print("[Внимание] Ответ ИИ не распознан как таблица — собираю чек-лист из отчёта локально.")
    return fallback_checklist_rows(report_text), answer


def save_checklist(rows: list[list[str]], raw_ai_answer: str = "") -> list[Path]:
    """
    Сохранение чек-листа отдельными файлами: CSV-таблица (для Excel) и
    markdown-таблица (для просмотра). Возвращает список созданных файлов.
    """
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime(TIMESTAMP_FORMAT)
    saved: list[Path] = []

    # 1) CSV: разделитель ';', UTF-8 BOM — корректно открывается в Excel.
    csv_path = OUTPUT_DIR / f"{CHECKLIST_PREFIX}{timestamp}{CHECKLIST_EXTENSION}"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(CHECKLIST_HEADER)
        writer.writerows(rows)
    saved.append(csv_path)

    # 2) Markdown-таблица с теми же столбцами.
    md_path = OUTPUT_DIR / f"{CHECKLIST_PREFIX}{timestamp}.md"
    lines = [
        "# Чек-лист проверок требований\n",
        f"- Дата и время формирования: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"- Проверок в чек-листе: {len(rows)}\n",
    ]
    if rows:
        lines.append("| " + " | ".join(CHECKLIST_HEADER) + " |")
        lines.append("|" + "---|" * len(CHECKLIST_HEADER))
        for r in rows:
            safe = [clean_md(c).replace("|", "\\|") for c in r]
            lines.append("| " + " | ".join(safe) + " |")
    else:
        lines.append("_Чек-лист пуст: не удалось извлечь проверки из ответа ИИ и отчёта._")
    if raw_ai_answer and not rows:
        lines.append("\n## Сырой ответ ИИ\n\n" + raw_ai_answer)
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    saved.append(md_path)
    return saved


# --------------------------------- main -------------------------------------

def main() -> int:
    print("=" * 70)
    print("Утилита тестирования требований с помощью ИИ")
    print("=" * 70)

    if not OPENAI_API_KEY:
        print("ОШИБКА: в .env не задан OPENAI_API_KEY (токен для доступа к ИИ).")
        return 1
    if not INPUT_DIR.exists():
        print(f"ОШИБКА: папка с требованиями не найдена: {INPUT_DIR}")
        return 1
    if not PROMPT_FILE.exists():
        print(f"ОШИБКА: файл промта не найден: {PROMPT_FILE}")
        return 1

    # Шаг 1. Сбор и чтение файлов требований.
    files = collect_requirement_files(INPUT_DIR)
    if not files:
        print(f"В папке {INPUT_DIR} не найдено файлов форматов: "
              f"{', '.join(SUPPORTED_EXTENSIONS)}")
        return 1

    print(f"\n[Шаг 1] Найдено файлов с требованиями: {len(files)}")
    for f in files:
        print(f"   - {f.relative_to(INPUT_DIR)} ({f.suffix})")

    requirements_text = build_requirements_payload(files, INPUT_DIR)
    prompt_text = PROMPT_FILE.read_text(encoding="utf-8").strip()
    print(f"\n[Шаг 1] Содержимое собрано: {len(requirements_text)} символов, "
          f"промт: {len(prompt_text)} символов.")

    # Шаг 2. Запрос к ИИ. Функция блокируется до получения полного ответа.
    print(f"\n[Шаг 2] Отправка данных в ИИ ({OPENAI_BASE_URL})...")
    try:
        answer = ask_ai(prompt_text, requirements_text)
    except Exception as e:
        print(f"\nОШИБКА при запросе к ИИ: {e}")
        return 1

    # Шаг 3. Обработка и сохранение результата — только после получения ответа.
    out_path = save_result(answer, len(files))
    print(f"\n[Шаг 3] Результат ИИ сохранён в файл: {out_path}")

    # Шаг 4. После завершения тестирования — mind-map и чек-лист (отдельные файлы).
    # Каждый подзапрос также дожидается полного ответа ИИ перед обработкой.
    if GENERATE_MINDMAP or GENERATE_CHECKLIST:
        print("\n[Шаг 4] Формирование дополнительных артефактов по результатам тестирования...")

    if GENERATE_MINDMAP:
        mindmap_text = build_mindmap(answer, requirements_text)
        mm_path = save_mindmap(mindmap_text)
        print(f"[Шаг 4] Mind-map сохранена в файл: {mm_path}")

    if GENERATE_CHECKLIST:
        rows, raw = generate_checklist(answer, requirements_text)
        if not rows:
            print("[Внимание] Чек-лист получился пустым — сохраняю только сырой ответ ИИ.")
        paths = save_checklist(rows, raw)
        for p in paths:
            print(f"[Шаг 4] Чек-лист сохранён в файл: {p}")

    print("\nГотово.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
